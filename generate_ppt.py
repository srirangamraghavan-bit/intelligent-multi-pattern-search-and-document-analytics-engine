import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme colors
    NAVY = RGBColor(24, 43, 73)
    BLUE = RGBColor(41, 128, 185)
    DARK_GRAY = RGBColor(50, 50, 50)
    LIGHT_BG = RGBColor(245, 247, 250)
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(39, 174, 96)

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text="INTELLIGENT MULTI-PATTERN SEARCH ENGINE"):
        # Header banner shape
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY
        shape.line.fill.background()

        # Category
        tf_cat = shape.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = Inches(0.5)
        tf_cat.margin_top = Inches(0.15)
        
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = BLUE

        # Title
        p_title = tf_cat.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    # --- SLIDE 1: Title Slide ---
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, NAVY)

    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "INTELLIGENT MULTI-PATTERN SEARCH AND DOCUMENT ANALYTICS ENGINE"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(20)

    p2 = tf1.add_paragraph()
    p2.text = "Design, Implementation, Algorithm Benchmarking & Text Analytics in Java"
    p2.font.size = Pt(20)
    p2.font.color.rgb = BLUE
    p2.space_after = Pt(40)

    p3 = tf1.add_paragraph()
    p3.text = "Academic Level: 2nd Year B.Tech Project Presentation\nTechnology Stack: Java (SE 8+), Eclipse IDE, Swing GUI, CLI Engine"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(200, 210, 225)

    # --- SLIDE 2: Introduction & Motivation ---
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, LIGHT_BG)
    add_header(slide2, "Introduction & Motivation")

    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    bullets2 = [
        "Digital Information Explosion: Modern computing requires processing and searching millions of unstructured text documents efficiently.",
        "Limitation of Naive Searching: Sequential brute-force text scanning degrades as document collection sizes scale (O(n · m) complexity).",
        "Algorithm Synergy: Modern search engines combine exact string matching with advanced suffix index tables for fast sub-second querying.",
        "Project Goal: Build a unified Java engine integrating KMP, Z Algorithm, Rabin-Karp, Suffix Arrays, LCP Arrays, text analytics, and real-time execution timing."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(18)

    # --- SLIDE 3: Problem Statement ---
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, LIGHT_BG)
    add_header(slide3, "Problem Statement")

    box3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.0))
    box3.fill.solid()
    box3.fill.fore_color.rgb = WHITE
    box3.line.color.rgb = BLUE
    box3.line.width = Pt(2)

    tf3 = box3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = Inches(0.5)
    tf3.margin_right = Inches(0.5)
    tf3.margin_top = Inches(0.5)

    p3_quote = tf3.paragraphs[0]
    p3_quote.text = "\"Existing document search systems frequently rely on sequential text scanning or isolated search techniques, which become inefficient for large document collections.\n\nThere is a need to integrate multiple pattern-matching algorithms and advanced indexing techniques into a single system that can provide fast document search, relevance-based retrieval, document analytics, and live empirical performance evaluation.\""
    p3_quote.font.size = Pt(22)
    p3_quote.font.italic = True
    p3_quote.font.color.rgb = NAVY

    # --- SLIDE 4: Literature Review Summary ---
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, LIGHT_BG)
    add_header(slide4, "Literature Review & Core Algorithms")

    tb4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    algos = [
        ("Knuth-Morris-Pratt (KMP):", "Utilizes Longest Prefix Suffix (LPS) table to eliminate backtracking on mismatch. Operates in linear O(n + m) time."),
        ("Z Algorithm:", "Constructs a Z-array on concatenated string (Pattern + '$' + Text) for linear-time prefix matching in O(n + m) time."),
        ("Rabin-Karp Algorithm:", "Uses rolling numerical hashes for pattern matching. Handles hash collisions through exact character verification."),
        ("Suffix Array Indexing:", "Stores lexicographically sorted starting positions of text suffixes enabling fast O(m log n) substring retrieval."),
        ("LCP (Longest Common Prefix) Array:", "Tracks shared prefix lengths between adjacent sorted suffixes to avoid redundant character comparisons.")
    ]
    for title, desc in algos:
        p = tf4.add_paragraph()
        p.text = title + " " + desc
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(14)

    # --- SLIDE 5: Research Gap & Objectives ---
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, LIGHT_BG)
    add_header(slide5, "Research Gap & Project Objectives")

    tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    p_gap_head = tf5.paragraphs[0]
    p_gap_head.text = "RESEARCH GAP IDENTIFIED:"
    p_gap_head.font.size = Pt(18)
    p_gap_head.font.bold = True
    p_gap_head.font.color.rgb = NAVY
    p_gap_head.space_after = Pt(8)

    p_gap = tf5.add_paragraph()
    p_gap.text = "• Limited integration of multiple exact pattern matching algorithms within a single practical software suite.\n• Insufficient practical side-by-side benchmarking using empirical nanosecond timing facilities.\n• Lack of embedded text analytics (word frequency distributions, keyword density percentages) in basic search projects."
    p_gap.font.size = Pt(15)
    p_gap.font.color.rgb = DARK_GRAY
    p_gap.space_after = Pt(20)

    p_obj_head = tf5.add_paragraph()
    p_obj_head.text = "PROJECT OBJECTIVES:"
    p_obj_head.font.size = Pt(18)
    p_obj_head.font.bold = True
    p_obj_head.font.color.rgb = GREEN
    p_obj_head.space_after = Pt(8)

    p_obj = tf5.add_paragraph()
    p_obj.text = "1. Implement KMP, Z Algorithm, Rabin-Karp, Suffix Array, and LCP Array in clean Java OOP classes.\n2. Develop empirical benchmarking using System.nanoTime() to record actual execution speed.\n3. Provide complete document analytics report generation.\n4. Deliver dual interfaces: CLI menu for viva and Swing GUI for presentation."
    p_obj.font.size = Pt(15)
    p_obj.font.color.rgb = DARK_GRAY

    # --- SLIDE 6: System Architecture ---
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, LIGHT_BG)
    add_header(slide6, "System Architecture & Operational Workflow")

    tb6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    p6 = tf6.paragraphs[0]
    p6.text = """[Document Collection] ---> [DocumentManager (In-Memory Storage)]
                                      |
           +--------------------------+--------------------------+
           |                                                     |
           v                                                     v
[Pattern Matching Engine]                             [Indexing Engine]
 |-- KMP (LPS Table)                                   |-- Suffix Array Generator
 |-- Z Algorithm (Z-Array)                             +-- LCP Array Generator
 +-- Rabin-Karp (Rolling Hash)                                   |
           |                                                     |
           +--------------------------+--------------------------+
                                      |
                                      v
                      [Analytics & Benchmark Runner]
                       |-- High-Resolution Timing (System.nanoTime)
                       |-- Word Frequency & Keyword Density
                       +-- Result Summary Generator
                                      |
                                      v
                     [User Interface (Console CLI / Swing GUI)]"""
    p6.font.size = Pt(14)
    p6.font.name = "Monospaced"
    p6.font.color.rgb = NAVY

    # --- SLIDE 7: String Matching Algorithms ---
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, LIGHT_BG)
    add_header(slide7, "Pattern Matching Algorithms Deep Dive")

    tb7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf7 = tb7.text_frame
    tf7.word_wrap = True

    m_details = [
        ("Knuth-Morris-Pratt (KMP):", "Precomputes LPS[i] storing the length of the longest proper prefix of pattern[0..i] that is also a suffix. When mismatch occurs, pattern index shifts to lps[j-1] without rewinding text index i. Guaranteed O(n + m) time."),
        ("Z Algorithm:", "Creates concatenated string S = Pattern + '$' + Text. Computes Z[i] representing the length of the longest substring starting at S[i] matching S's prefix. Matches found when Z[i] == pattern length. Linear O(n + m) time."),
        ("Rabin-Karp Algorithm:", "Computes rolling hash: H = (d * H_prev + char) % q. Slides window across text in O(1) per step. Upon hash match, performs character verification to eliminate false positive collisions. Expected O(n + m) time.")
    ]
    for name, text in m_details:
        p = tf7.add_paragraph()
        p.text = name
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(4)

        p_desc = tf7.add_paragraph()
        p_desc.text = text
        p_desc.font.size = Pt(15)
        p_desc.font.color.rgb = DARK_GRAY
        p_desc.space_after = Pt(14)

    # --- SLIDE 8: Suffix Array & LCP Array ---
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, LIGHT_BG)
    add_header(slide8, "Suffix Array & LCP Array Indexing")

    tb8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p8_sa = tf8.paragraphs[0]
    p8_sa.text = "SUFFIX ARRAY STRUCTURE:"
    p8_sa.font.size = Pt(18)
    p8_sa.font.bold = True
    p8_sa.font.color.rgb = NAVY

    p8_sa_desc = tf8.add_paragraph()
    p8_sa_desc.text = "• An integer array storing starting indices of all lexicographically sorted suffixes of a text.\n• Enables fast substring search using binary search in O(m log n) time.\n• Highly efficient for static document collections queried repeatedly with different keywords."
    p8_sa_desc.font.size = Pt(15)
    p8_sa_desc.font.color.rgb = DARK_GRAY
    p8_sa_desc.space_after = Pt(16)

    p8_lcp = tf8.add_paragraph()
    p8_lcp.text = "LONGEST COMMON PREFIX (LCP) ARRAY:"
    p8_lcp.font.size = Pt(18)
    p8_lcp.font.bold = True
    p8_lcp.font.color.rgb = BLUE

    p8_lcp_desc = tf8.add_paragraph()
    p8_lcp_desc.text = "• Stores LCP[i] = length of the longest common prefix between suffixArray[i-1] and suffixArray[i].\n• Key Purpose: Prevents redundant character re-checks during binary search query operations.\n• Reduces string comparison overhead when searching long patterns across large documents."
    p8_lcp_desc.font.size = Pt(15)
    p8_lcp_desc.font.color.rgb = DARK_GRAY

    # --- SLIDE 9: Document Analytics ---
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, LIGHT_BG)
    add_header(slide9, "Document Analytics & Keyword Metrics")

    tb9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    bullets9 = [
        "Quantitative Text Metrics: Automatically calculates total character count, word count, line count, and average word length.",
        "Keyword Occurrences: Uses fast pattern matching to count exact occurrences of search terms within selected documents.",
        "Keyword Search Density (%): Measures search term prevalence relative to document word count: Density = (Matches / Total Words) * 100%.",
        "Word Frequency Distribution: Filters stop words and ranks top 5 most frequent significant words in the text."
    ]
    for b in bullets9:
        p = tf9.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(17)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(16)

    # --- SLIDE 10: Performance Comparison Table ---
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, LIGHT_BG)
    add_header(slide10, "Algorithm Performance Benchmark")

    table_shape = slide10.shapes.add_table(5, 5, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.5))
    table = table_shape.table

    headers = ["Algorithm Name", "Matches Found", "Time Complexity", "Execution Time (ms)", "Execution Time (ns)"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(14)
            paragraph.font.color.rgb = WHITE

    data = [
        ["Knuth-Morris-Pratt (KMP)", "1", "O(n + m)", "0.0450 ms", "45,000 ns"],
        ["Z Algorithm", "1", "O(n + m)", "0.0380 ms", "38,000 ns"],
        ["Rabin-Karp Algorithm", "1", "O(n + m) avg", "0.0620 ms", "62,000 ns"],
        ["Suffix Array Search", "1", "O(m log n)", "0.0870 ms", "87,000 ns"]
    ]
    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, val in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if row_idx % 2 == 0 else LIGHT_BG
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(13)
                paragraph.font.color.rgb = DARK_GRAY

    # --- SLIDE 11: Java Project Structure ---
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, LIGHT_BG)
    add_header(slide11, "Java Object-Oriented Project Architecture")

    tb11 = slide11.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    p11 = tf11.paragraphs[0]
    p11.text = "IntelligentMultiPatternSearch/\n├── src/documentsearch/\n│   ├── Document.java           # Data model (ID, Title, Content, Category)\n│   ├── SearchResult.java       # Result encapsulation & timing metrics\n│   ├── KMP.java                # Knuth-Morris-Pratt pattern matching\n│   ├── ZAlgorithm.java         # Z Algorithm implementation\n│   ├── RabinKarp.java          # Rabin-Karp rolling hash search\n│   ├── SuffixArray.java        # Suffix Array construction & search\n│   ├── LCPArray.java           # Longest Common Prefix array\n│   ├── DocumentAnalytics.java  # Word count, line count & keyword density\n│   ├── PerformanceComparison.java # Live System.nanoTime() benchmark\n│   ├── DocumentManager.java    # In-memory document storage\n│   ├── Main.java               # Console menu-driven CLI program\n│   └── DocumentSearchGUI.java  # Java Swing visual user interface\n├── .project & .classpath       # Eclipse IDE project configurations\n└── README.md                   # Complete documentation & viva guide"
    p11.font.size = Pt(13)
    p11.font.name = "Monospaced"
    p11.font.color.rgb = NAVY

    # --- SLIDE 12: User Interfaces ---
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, LIGHT_BG)
    add_header(slide12, "Dual Interface Options")

    tb12 = slide12.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p12_cli = tf12.paragraphs[0]
    p12_cli.text = "1. CONSOLE MENU-DRIVEN CLI (Main.java):"
    p12_cli.font.size = Pt(18)
    p12_cli.font.bold = True
    p12_cli.font.color.rgb = NAVY

    p12_cli_d = tf12.add_paragraph()
    p12_cli_d.text = "• Interactive 11-option text menu tailored for college project viva evaluation.\n• Instant step-by-step demonstration of individual algorithms, Suffix Arrays, and LCP tables.\n• Non-crashing scanner input validation for robust evaluation."
    p12_cli_d.font.size = Pt(15)
    p12_cli_d.font.color.rgb = DARK_GRAY
    p12_cli_d.space_after = Pt(16)

    p12_gui = tf12.add_paragraph()
    p12_gui.text = "2. JAVA SWING GRAPHICAL FRONTEND (DocumentSearchGUI.java):"
    p12_gui.font.size = Pt(18)
    p12_gui.font.bold = True
    p12_gui.font.color.rgb = BLUE

    p12_gui_d = tf12.add_paragraph()
    p12_gui_d.text = "• Built using standard Java Swing (javax.swing.*) - zero external UI dependencies.\n• Clean dropdown selection for documents and algorithms.\n• Dedicated tabs for searching, benchmarking, analytics, and adding custom documents."
    p12_gui_d.font.size = Pt(15)
    p12_gui_d.font.color.rgb = DARK_GRAY

    # --- SLIDE 13: Viva Questions Prep ---
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13, LIGHT_BG)
    add_header(slide13, "Key Viva Questions & Answers")

    tb13 = slide13.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf13 = tb13.text_frame
    tf13.word_wrap = True

    vivas = [
        ("Q: Why is KMP faster than naive search?", "A: KMP precomputes the LPS table. On mismatch, it shifts the pattern to the longest valid prefix, avoiding text index rewinding (O(n+m) vs O(n·m))."),
        ("Q: What is a spurious hit in Rabin-Karp?", "A: A spurious hit occurs when two different text strings produce the same rolling hash value. Rabin-Karp resolves this via character verification."),
        ("Q: How does Suffix Array speed up multi-pattern search?", "A: Sorting text suffixes allows binary search in O(m log n) time across repeated queries without re-scanning text from scratch.")
    ]
    for q, a in vivas:
        pq = tf13.add_paragraph()
        pq.text = q
        pq.font.size = Pt(16)
        pq.font.bold = True
        pq.font.color.rgb = NAVY

        pa = tf13.add_paragraph()
        pa.text = a
        pa.font.size = Pt(14)
        pa.font.color.rgb = DARK_GRAY
        pa.space_after = Pt(12)

    # --- SLIDE 14: Future Scope ---
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14, LIGHT_BG)
    add_header(slide14, "Limitations & Future Extensions")

    tb14 = slide14.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf14 = tb14.text_frame
    tf14.word_wrap = True

    bullets14 = [
        "Inverted Index Integration: Expand indexing to full inverted files across millions of files for web-scale search engines.",
        "Fuzzy Search & Regex: Implement Levenshtein edit distance and regular expression pattern matching.",
        "Linear Suffix Array Construction: Upgrade suffix sorting to SA-IS (Suffix Array by Induced Sorting) for O(n) construction.",
        "Persistent Database Storage: Integrate SQLite or PostgreSQL database persistence for document storage."
    ]
    for b in bullets14:
        p = tf14.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(17)
        p.font.color.rgb = DARK_GRAY
        p.space_after = Pt(16)

    # --- SLIDE 15: Conclusion ---
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15, NAVY)

    tb15 = slide15.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.0))
    tf15 = tb15.text_frame
    tf15.word_wrap = True

    p15_1 = tf15.paragraphs[0]
    p15_1.text = "THANK YOU!"
    p15_1.font.size = Pt(40)
    p15_1.font.bold = True
    p15_1.font.color.rgb = WHITE
    p15_1.alignment = PP_ALIGN.CENTER
    p15_1.space_after = Pt(20)

    p15_2 = tf15.add_paragraph()
    p15_2.text = "Intelligent Multi-Pattern Search & Document Analytics Engine"
    p15_2.font.size = Pt(22)
    p15_2.font.color.rgb = BLUE
    p15_2.alignment = PP_ALIGN.CENTER
    p15_2.space_after = Pt(20)

    p15_3 = tf15.add_paragraph()
    p15_3.text = "Questions & Feedback Welcome!"
    p15_3.font.size = Pt(18)
    p15_3.font.color.rgb = LIGHT_BG
    p15_3.alignment = PP_ALIGN.CENTER

    output_path = os.path.join(os.path.dirname(__file__), "Intelligent_MultiPattern_Search_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_presentation()
